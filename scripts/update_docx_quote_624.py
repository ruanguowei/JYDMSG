from copy import deepcopy
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_or_replace_total_row(table, total_text):
    # Remove an existing total row if the script is run again.
    for row in list(table.rows):
        if "合计" in " ".join(cell.text for cell in row.cells):
            table._tbl.remove(row._tr)

    row = table.add_row()
    for cell in row.cells:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    merged = row.cells[0].merge(row.cells[1])
    set_cell_text(merged, "合计", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_cell_text(row.cells[2], total_text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

    for cell in row.cells:
        set_cell_shading(cell, "D9EAF7")


def main():
    desktop = Path.home() / "Desktop"
    matches = sorted(desktop.glob("*6.24.docx"))
    if not matches:
        raise FileNotFoundError("未找到桌面上的 6.24 docx 文件")
    docx_path = matches[0]
    backup_path = docx_path.with_name(docx_path.stem + "_报价前备份" + docx_path.suffix)

    if not backup_path.exists():
        backup_path.write_bytes(docx_path.read_bytes())

    doc = Document(str(docx_path))
    table = doc.tables[0]

    prices = [
        (2, "￥6,000"),
        (3, "￥5,000"),
        (4, "￥3,000"),
        (5, "￥6,000"),
        (6, "￥6,500"),
        (7, "￥7,000"),
        (8, "￥4,500"),
        (9, "￥3,800"),
        (10, "￥5,000"),
        (11, "￥3,000"),
    ]

    # Update header amount wording.
    header = table.cell(1, 1)
    header.text = re.sub(r"费用\s*50000\s*元", "费用49800元", header.text)
    for paragraph in header.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.name = "宋体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(10.5)

    for row_idx, amount in prices:
        set_cell_text(table.cell(row_idx, 2), amount, bold=False)

    add_or_replace_total_row(table, "￥49,800")

    doc.save(str(docx_path))
    print(docx_path)
    print(backup_path)


if __name__ == "__main__":
    main()
